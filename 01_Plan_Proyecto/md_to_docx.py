import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

SRC = Path(r"C:\Users\User\Downloads\Proyecto_Desarrollo\01_Plan_Proyecto\Plan_Proyecto_TUPA_UNSAAC.md")
OUT = SRC.with_suffix(".docx")

def set_cell_bg(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {
        qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): color_hex,
    })
    tcPr.append(shd)

def add_inline(paragraph, text):
    # handle **bold** and *italic*
    tokens = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = paragraph.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("*") and tok.endswith("*"):
            r = paragraph.add_run(tok[1:-1]); r.italic = True
        else:
            paragraph.add_run(tok)

def parse_table(lines, i):
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip().strip("|")
        cells = [c.strip() for c in raw.split("|")]
        if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            rows.append(cells)
        i += 1
    return rows, i

def main():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.5); section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5); section.right_margin = Cm(2.5)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"; style.font.size = Pt(11)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    first_h1 = True
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped or stripped == "---":
            i += 1; continue

        m = re.match(r"^(#{1,6})\s+(.*)", stripped)
        if m:
            level = len(m.group(1)); text = m.group(2)
            if level == 1 and first_h1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(text.upper()); r.bold = True; r.font.size = Pt(18)
                first_h1 = False
            else:
                h = doc.add_heading(level=min(level, 4))
                add_inline(h, text)
            i += 1; continue

        if stripped.startswith("|"):
            rows, i = parse_table(lines, i)
            if not rows: continue
            ncol = max(len(r) for r in rows)
            table = doc.add_table(rows=0, cols=ncol)
            table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(rows):
                cells = table.add_row().cells
                for ci in range(ncol):
                    txt = row[ci] if ci < len(row) else ""
                    para = cells[ci].paragraphs[0]
                    add_inline(para, txt)
                    for run in para.runs:
                        run.font.size = Pt(10)
                        if ri == 0: run.bold = True
                    if ri == 0: set_cell_bg(cells[ci], "D9E2F3")
            doc.add_paragraph()
            continue

        if re.match(r"^[-*]\s+", stripped):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^[-*]\s+", "", stripped))
            i += 1; continue

        if re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s+", "", stripped))
            i += 1; continue

        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            add_inline(p, stripped.lstrip("> ").strip())
            for run in p.runs: run.italic = True
            i += 1; continue

        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1

    doc.save(OUT)
    print(f"OK -> {OUT}")

if __name__ == "__main__":
    main()
